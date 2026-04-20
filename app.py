import json
import os
import sys
import signal
import webbrowser
import threading
import tempfile
import time
from typing import Optional

from flask import Flask, request, send_file, jsonify, render_template
from openpyxl import load_workbook
import pandas as pd

from extract import extract_all, run as run_extraction


# ---------------------------------------------------------------------------
# App setup
# ---------------------------------------------------------------------------

def _get_template_dir():
    if getattr(sys, "frozen", False):
        return os.path.join(sys._MEIPASS, "templates")
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates")


app = Flask(__name__, template_folder=_get_template_dir())

UPLOAD_DIR = tempfile.mkdtemp(prefix="data_sorter_")

if getattr(sys, "frozen", False):
    HERE = os.path.dirname(sys.executable)
else:
    HERE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(HERE, "data")
JSON_PATH = os.path.join(HERE, "performance_data.json")

_data = {}


def _load_data():
    global _data
    os.makedirs(DATA_DIR, exist_ok=True)
    if not os.path.exists(JSON_PATH):
        run_extraction(DATA_DIR, JSON_PATH)
    if os.path.exists(JSON_PATH):
        with open(JSON_PATH, "r", encoding="utf-8") as f:
            _data = json.load(f)


_load_data()

# ---------------------------------------------------------------------------
# Heartbeat / watchdog
# ---------------------------------------------------------------------------
_last_heartbeat = time.time()
_HEARTBEAT_TIMEOUT = 120


def _watchdog():
    while True:
        time.sleep(10)
        if time.time() - _last_heartbeat > _HEARTBEAT_TIMEOUT:
            os.kill(os.getpid(), signal.SIGTERM)
            return


@app.route("/api/heartbeat", methods=["POST"])
def heartbeat():
    global _last_heartbeat
    _last_heartbeat = time.time()
    return "", 204


# ---------------------------------------------------------------------------
# Page routes
# ---------------------------------------------------------------------------

@app.route("/")
def page_processor():
    return render_template("processor.html")


@app.route("/dashboard")
def page_dashboard():
    return render_template("dashboard.html")


# ---------------------------------------------------------------------------
# Processor helper: save an uploaded FileStorage to a temp file
# ---------------------------------------------------------------------------

_save_counter = 0


def _save_upload(file_storage, prefix="upload"):
    global _save_counter
    _save_counter += 1
    path = os.path.join(UPLOAD_DIR, f"{prefix}_{_save_counter}.xlsx")
    file_storage.save(path)
    return path


# ---------------------------------------------------------------------------
# Processor: Department Evaluation -> 部门绩效考核, row 6
# ---------------------------------------------------------------------------

def _read_dept_eval(file_path):
    wb = load_workbook(file_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=2, max_row=ws.max_row,
                             min_col=2, max_col=13, values_only=True))
    wb.close()
    return pd.DataFrame(rows)


def _calc_dept_score(col):
    col = col.dropna().astype(str).str.strip().str.upper()
    if "C" in col.values:
        return "C"
    if (col == "A").sum() > 8:
        return "A"
    return "B"


@app.route("/api/process/department", methods=["POST"])
def process_department():
    target = request.files.get("target")
    file1 = request.files.get("file_regular")
    file2 = request.files.get("file_leader")
    if not target or not file1 or not file2:
        return jsonify(error="请上传目标文件、部门考核评定表和领导评定表"), 400

    path1 = _save_upload(file1, "dept_regular")
    path2 = _save_upload(file2, "dept_leader")

    df1 = _read_dept_eval(path1)
    df2 = _read_dept_eval(path2)
    combined = pd.concat([df1, df2], ignore_index=True)
    scores = combined.apply(_calc_dept_score).tolist()

    target_path = os.path.join(UPLOAD_DIR, "target_dept.xlsx")
    target.save(target_path)
    wb = load_workbook(target_path)
    ws = wb["部门绩效考核"]
    for i, score in enumerate(scores[:12]):
        ws.cell(row=6, column=2 + i).value = score
    wb.save(target_path)
    wb.close()

    return send_file(target_path, as_attachment=True,
                     download_name=request.form.get("download_name", "output.xlsx"))


# ---------------------------------------------------------------------------
# Processor: Middle Management -> 部门副职考核, row 12
# ---------------------------------------------------------------------------

DEPT_NAME_MAP = {
    "物业分公司": "物业分公司",
    "高新区分公司": "高新分公司",
    "综合办公室": "综合办公室",
    "投资拓展部": "投资拓展部",
    "项目管理部": "项目管理部",
    "造价合约部": "造价合约部",
}


def _extract_dept_from_header(header: str) -> Optional[str]:
    clean = header.replace("（必填）", "").replace("（已删除）", "").strip()
    parts = clean.split()
    return parts[0] if parts else None


def _read_mgmt_scores(target_path, num_cols=6):
    wb_data = load_workbook(target_path, data_only=True)
    ws_data = wb_data["部门副职考核"]
    scores = []
    for col in range(2, 2 + num_cols):
        scores.append(ws_data.cell(row=9, column=col).value)
    wb_data.close()
    return scores


def _compute_mgmt_ranks(scores):
    numeric = []
    for s in scores:
        try:
            numeric.append(float(s))
        except (TypeError, ValueError):
            return None

    sorted_desc = sorted(set(numeric), reverse=True)
    rank_map = {}
    pos = 1
    for val in sorted_desc:
        if val not in rank_map:
            rank_map[val] = pos
        pos += sum(1 for v in numeric if v == val)

    return [rank_map[s] for s in numeric]


@app.route("/api/process/middlemgmt", methods=["POST"])
def process_middle_mgmt():
    target = request.files.get("target")
    source = request.files.get("file_mgmt")
    if not target or not source:
        return jsonify(error="请上传目标文件和中层副职考核评定表"), 400

    src_path = _save_upload(source, "mgmt")
    wb_src = load_workbook(src_path, data_only=True)
    ws_src = wb_src[wb_src.sheetnames[0]]

    headers = [cell.value for cell in ws_src[1]]
    src_dept_col = {}
    for col_idx in range(1, len(headers)):
        h = headers[col_idx]
        if h is None:
            continue
        dept = _extract_dept_from_header(str(h))
        if dept and dept in DEPT_NAME_MAP:
            votes = []
            for row in range(2, ws_src.max_row + 1):
                v = ws_src.cell(row=row, column=col_idx + 1).value
                if v is not None:
                    votes.append(str(v).strip().upper())
            src_dept_col[dept] = votes
    wb_src.close()

    target_path = os.path.join(UPLOAD_DIR, "target_mgmt.xlsx")
    target.save(target_path)

    mgmt_scores = _read_mgmt_scores(target_path)
    mgmt_ranks = _compute_mgmt_ranks(mgmt_scores)

    wb = load_workbook(target_path)
    ws = wb["部门副职考核"]

    target_depts = []
    for col in range(2, ws.max_column + 1):
        v = ws.cell(row=3, column=col).value
        if v:
            target_depts.append((col, str(v).strip()))

    for col_idx, dept_name in target_depts:
        src_key = None
        for k, mapped in DEPT_NAME_MAP.items():
            if mapped == dept_name:
                src_key = k
                break
        if src_key and src_key in src_dept_col:
            votes = src_dept_col[src_key]
            a_count = sum(1 for v in votes if v == "A")
            result = "A" if a_count > 8 else "B"
            ws.cell(row=12, column=col_idx).value = result

    if mgmt_ranks:
        for i, rank in enumerate(mgmt_ranks):
            ws.cell(row=11, column=2 + i).value = rank

    wb.save(target_path)
    wb.close()

    return send_file(target_path, as_attachment=True,
                     download_name=request.form.get("download_name", "output.xlsx"))


# ---------------------------------------------------------------------------
# Processor: Employee Rating -> 绩效考评汇总表
# ---------------------------------------------------------------------------

def _competition_rank(scores):
    ranks = []
    for i, s in enumerate(scores):
        if i == 0 or s != scores[i - 1]:
            ranks.append(i + 1)
        else:
            ranks.append(ranks[-1])
    return ranks


def _find_department_sections(ws):
    header_row = 3
    dept_starts = []

    for row in range(header_row + 1, ws.max_row + 1):
        val_a = ws.cell(row=row, column=1).value
        if val_a and str(val_a).strip():
            dept_name = str(val_a).strip()
            if dept_name.startswith("备注"):
                break
            dept_starts.append((row, dept_name))

    sections = []
    for idx, (start_row, dept_name) in enumerate(dept_starts):
        if idx + 1 < len(dept_starts):
            end_row = dept_starts[idx + 1][0] - 1
        else:
            for r in range(start_row, ws.max_row + 1):
                val = ws.cell(row=r, column=1).value
                if val and str(val).strip().startswith("备注"):
                    end_row = r - 1
                    break
            else:
                end_row = ws.max_row
        sections.append((dept_name, start_row, end_row))

    return sections


@app.route("/api/process/employee", methods=["POST"])
def process_employee():
    target = request.files.get("target")
    source = request.files.get("file_employee")
    if not target or not source:
        return jsonify(error="请上传目标文件和员工等级评定表"), 400

    src_path = _save_upload(source, "employee")
    wb_src = load_workbook(src_path, data_only=True)
    ws_src = wb_src[wb_src.sheetnames[0]]
    employees = []
    for row in range(2, ws_src.max_row + 1):
        name = ws_src.cell(row=row, column=1).value
        dept = ws_src.cell(row=row, column=2).value
        score = ws_src.cell(row=row, column=3).value
        grade = ws_src.cell(row=row, column=4).value
        if name and dept and score is not None:
            employees.append({
                "name": str(name).strip(),
                "dept": str(dept).strip(),
                "score": float(str(score)),
                "grade": str(grade).strip() if grade else "",
            })
    wb_src.close()

    dept_groups = {}
    for emp in employees:
        dept_groups.setdefault(emp["dept"], []).append(emp)
    for dept in dept_groups:
        dept_groups[dept].sort(key=lambda e: e["score"], reverse=True)
        scores = [e["score"] for e in dept_groups[dept]]
        ranks = _competition_rank(scores)
        for emp, rank in zip(dept_groups[dept], ranks):
            emp["rank"] = rank

    target_path = os.path.join(UPLOAD_DIR, "target_emp.xlsx")
    target.save(target_path)
    wb = load_workbook(target_path)
    ws = wb["绩效考评汇总表"]

    sections = _find_department_sections(ws)

    for dept_name, start_row, end_row in sections:
        matched = dept_groups.get(dept_name)
        if not matched:
            continue
        available_rows = end_row - start_row + 1
        for i, emp in enumerate(matched[:available_rows]):
            row = start_row + i
            ws.cell(row=row, column=7).value = emp["name"]
            ws.cell(row=row, column=8).value = emp["rank"]
            ws.cell(row=row, column=9).value = emp["score"]
            ws.cell(row=row, column=10).value = emp["grade"]

    wb.save(target_path)
    wb.close()

    return send_file(target_path, as_attachment=True,
                     download_name=request.form.get("download_name", "output.xlsx"))


# ---------------------------------------------------------------------------
# Processor: Combined processing (all three at once)
# ---------------------------------------------------------------------------

@app.route("/api/process/all", methods=["POST"])
def process_all():
    target = request.files.get("target")
    file_regular = request.files.get("file_regular")
    file_leader = request.files.get("file_leader")
    file_mgmt = request.files.get("file_mgmt")
    file_employee = request.files.get("file_employee")

    if not target:
        return jsonify(error="请上传目标文件"), 400

    target_path = os.path.join(UPLOAD_DIR, "target_all.xlsx")
    target.save(target_path)

    mgmt_scores = _read_mgmt_scores(target_path)
    mgmt_ranks = _compute_mgmt_ranks(mgmt_scores)

    if file_regular and file_leader:
        path1 = _save_upload(file_regular, "all_dept_reg")
        path2 = _save_upload(file_leader, "all_dept_ldr")
        df1 = _read_dept_eval(path1)
        df2 = _read_dept_eval(path2)
        combined = pd.concat([df1, df2], ignore_index=True)
        scores = combined.apply(_calc_dept_score).tolist()

        wb = load_workbook(target_path)
        ws = wb["部门绩效考核"]
        for i, score in enumerate(scores[:12]):
            ws.cell(row=6, column=2 + i).value = score
        wb.save(target_path)
        wb.close()

    if file_mgmt:
        mgmt_path = _save_upload(file_mgmt, "all_mgmt")
        wb_src = load_workbook(mgmt_path, data_only=True)
        ws_src = wb_src[wb_src.sheetnames[0]]
        headers = [cell.value for cell in ws_src[1]]
        src_dept_col = {}
        for col_idx in range(1, len(headers)):
            h = headers[col_idx]
            if h is None:
                continue
            dept = _extract_dept_from_header(str(h))
            if dept and dept in DEPT_NAME_MAP:
                votes = []
                for row in range(2, ws_src.max_row + 1):
                    v = ws_src.cell(row=row, column=col_idx + 1).value
                    if v is not None:
                        votes.append(str(v).strip().upper())
                src_dept_col[dept] = votes
        wb_src.close()

        wb = load_workbook(target_path)
        ws = wb["部门副职考核"]
        target_depts = []
        for col in range(2, ws.max_column + 1):
            v = ws.cell(row=3, column=col).value
            if v:
                target_depts.append((col, str(v).strip()))
        for col_idx, dept_name in target_depts:
            src_key = None
            for k, mapped in DEPT_NAME_MAP.items():
                if mapped == dept_name:
                    src_key = k
                    break
            if src_key and src_key in src_dept_col:
                votes = src_dept_col[src_key]
                a_count = sum(1 for v in votes if v == "A")
                result = "A" if a_count > 8 else "B"
                ws.cell(row=12, column=col_idx).value = result
        if mgmt_ranks:
            for i, rank in enumerate(mgmt_ranks):
                ws.cell(row=11, column=2 + i).value = rank
        wb.save(target_path)
        wb.close()
    elif mgmt_ranks:
        wb = load_workbook(target_path)
        ws = wb["部门副职考核"]
        for i, rank in enumerate(mgmt_ranks):
            ws.cell(row=11, column=2 + i).value = rank
        wb.save(target_path)
        wb.close()

    if file_employee:
        emp_path = _save_upload(file_employee, "all_emp")
        wb_src = load_workbook(emp_path, data_only=True)
        ws_src = wb_src[wb_src.sheetnames[0]]
        employees = []
        for row in range(2, ws_src.max_row + 1):
            name = ws_src.cell(row=row, column=1).value
            dept = ws_src.cell(row=row, column=2).value
            score = ws_src.cell(row=row, column=3).value
            grade = ws_src.cell(row=row, column=4).value
            if name and dept and score is not None:
                employees.append({
                    "name": str(name).strip(),
                    "dept": str(dept).strip(),
                    "score": float(str(score)),
                    "grade": str(grade).strip() if grade else "",
                })
        wb_src.close()

        dept_groups = {}
        for emp in employees:
            dept_groups.setdefault(emp["dept"], []).append(emp)
        for dept in dept_groups:
            dept_groups[dept].sort(key=lambda e: e["score"], reverse=True)
            scores_list = [e["score"] for e in dept_groups[dept]]
            ranks = _competition_rank(scores_list)
            for emp, rank in zip(dept_groups[dept], ranks):
                emp["rank"] = rank

        wb = load_workbook(target_path)
        ws = wb["绩效考评汇总表"]
        sections = _find_department_sections(ws)
        for dept_name, start_row, end_row in sections:
            matched = dept_groups.get(dept_name)
            if not matched:
                continue
            available_rows = end_row - start_row + 1
            for i, emp in enumerate(matched[:available_rows]):
                row = start_row + i
                ws.cell(row=row, column=7).value = emp["name"]
                ws.cell(row=row, column=8).value = emp["rank"]
                ws.cell(row=row, column=9).value = emp["score"]
                ws.cell(row=row, column=10).value = emp["grade"]
        wb.save(target_path)
        wb.close()

    return send_file(target_path, as_attachment=True,
                     download_name=request.form.get("download_name", "output.xlsx"))


# ---------------------------------------------------------------------------
# Dashboard helpers
# ---------------------------------------------------------------------------

def _ym_key(year, month):
    return year * 100 + month


def _parse_ym(s):
    parts = s.split("-")
    return int(parts[0]), int(parts[1])


def _in_range(year, month, start, end):
    k = _ym_key(year, month)
    return _ym_key(*start) <= k <= _ym_key(*end)


# ---------------------------------------------------------------------------
# Dashboard API endpoints
# ---------------------------------------------------------------------------

@app.route("/api/time-range")
def api_time_range():
    all_ym = set()
    for recs in _data.get("people", {}).values():
        for r in recs:
            all_ym.add((r["year"], r["month"]))
    for recs in _data.get("office_scores", {}).values():
        for r in recs:
            all_ym.add((r["year"], r["month"]))

    if not all_ym:
        return jsonify({"months": []})

    sorted_ym = sorted(all_ym)
    return jsonify({
        "months": [{"year": y, "month": m} for y, m in sorted_ym],
    })


@app.route("/api/offices")
def api_offices():
    start_s = request.args.get("start")
    end_s = request.args.get("end")

    offices = set()

    if start_s and end_s:
        start, end = _parse_ym(start_s), _parse_ym(end_s)
        for name, recs in _data.get("people", {}).items():
            for r in recs:
                if _in_range(r["year"], r["month"], start, end):
                    offices.add(r["office"])
        for office, recs in _data.get("office_scores", {}).items():
            for r in recs:
                if _in_range(r["year"], r["month"], start, end):
                    offices.add(office)
    else:
        for recs in _data.get("people", {}).values():
            for r in recs:
                offices.add(r["office"])
        for office in _data.get("office_scores", {}):
            offices.add(office)

    return jsonify({"offices": sorted(offices)})


@app.route("/api/people")
def api_people():
    office = request.args.get("office", "")
    start_s = request.args.get("start")
    end_s = request.args.get("end")

    results = []

    for name, recs in _data.get("people", {}).items():
        matched = False
        ptype = "employee"
        for r in recs:
            in_time = True
            if start_s and end_s:
                start, end = _parse_ym(start_s), _parse_ym(end_s)
                in_time = _in_range(r["year"], r["month"], start, end)
            if in_time and r["office"] == office:
                matched = True
                ptype = r["type"]
        if matched:
            results.append({"name": name, "type": ptype})

    results.sort(key=lambda x: (0 if x["type"] == "manager" else 1, x["name"]))
    return jsonify({"people": results})


@app.route("/api/person/<name>")
def api_person(name):
    start_s = request.args.get("start")
    end_s = request.args.get("end")

    recs = _data.get("people", {}).get(name, [])
    if not recs:
        return jsonify({"error": "Person not found"}), 404

    filtered = []
    for r in recs:
        if start_s and end_s:
            start, end = _parse_ym(start_s), _parse_ym(end_s)
            if not _in_range(r["year"], r["month"], start, end):
                continue
        filtered.append(r)

    filtered.sort(key=lambda r: _ym_key(r["year"], r["month"]))

    labels = []
    marks = []
    result_counts = {"A": 0, "B": 0, "C": 0}
    details = []

    for r in filtered:
        label = f"{r['year']}-{r['month']:02d}"
        labels.append(label)
        marks.append(r["mark"])
        if r["result"] in result_counts:
            result_counts[r["result"]] += 1
        details.append({
            "month": label,
            "office": r["office"],
            "type": r["type"],
            "mark": r["mark"],
            "result": r["result"],
        })

    return jsonify({
        "name": name,
        "labels": labels,
        "marks": marks,
        "result_counts": result_counts,
        "details": details,
    })


@app.route("/api/office/<name>")
def api_office(name):
    start_s = request.args.get("start")
    end_s = request.args.get("end")

    recs = _data.get("office_scores", {}).get(name, [])
    if not recs:
        return jsonify({"error": "Office not found"}), 404

    filtered = []
    for r in recs:
        if start_s and end_s:
            start, end = _parse_ym(start_s), _parse_ym(end_s)
            if not _in_range(r["year"], r["month"], start, end):
                continue
        filtered.append(r)

    filtered.sort(key=lambda r: _ym_key(r["year"], r["month"]))

    labels = []
    marks = []
    result_counts = {"A": 0, "B": 0, "C": 0}
    details = []

    for r in filtered:
        label = f"{r['year']}-{r['month']:02d}"
        labels.append(label)
        marks.append(r["mark"])
        if r["result"] in result_counts:
            result_counts[r["result"]] += 1
        details.append({
            "month": label,
            "office": name,
            "type": "department",
            "mark": r["mark"],
            "result": r["result"],
        })

    return jsonify({
        "name": name,
        "labels": labels,
        "marks": marks,
        "result_counts": result_counts,
        "details": details,
    })


@app.route("/api/upload", methods=["POST"])
def api_upload():
    files = request.files.getlist("files")
    if not files or all(f.filename == "" for f in files):
        return jsonify({"ok": False, "error": "No files provided"}), 400

    uploaded = []
    for f in files:
        if not f.filename:
            continue
        fname = os.path.basename(f.filename)
        dest = os.path.join(DATA_DIR, fname)
        f.save(dest)
        uploaded.append(fname)

    try:
        run_extraction(DATA_DIR, JSON_PATH)
        _load_data()
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500

    return jsonify({
        "ok": True,
        "uploaded": uploaded,
        "files_count": len(_data.get("files_processed", [])),
        "people_count": len(_data.get("people", {})),
    })


@app.route("/api/refresh", methods=["POST"])
def api_refresh():
    try:
        run_extraction(DATA_DIR, JSON_PATH)
        _load_data()
        return jsonify({
            "ok": True,
            "files": _data.get("files_processed", []),
            "people_count": len(_data.get("people", {})),
        })
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500


@app.route("/api/clear", methods=["POST"])
def api_clear():
    global _data
    removed = []
    try:
        for fname in os.listdir(DATA_DIR):
            if fname.startswith("~$"):
                continue
            if fname.endswith(".xlsx"):
                os.remove(os.path.join(DATA_DIR, fname))
                removed.append(fname)
        if os.path.exists(JSON_PATH):
            os.remove(JSON_PATH)
        _data = {
            "last_updated": None,
            "files_processed": [],
            "office_scores": {},
            "people": {},
        }
        return jsonify({"ok": True, "removed": len(removed)})
    except Exception as exc:
        return jsonify({"ok": False, "error": str(exc)}), 500


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def _open_browser():
    webbrowser.open("http://127.0.0.1:5000")


if __name__ == "__main__":
    wd = threading.Thread(target=_watchdog, daemon=True)
    wd.start()

    threading.Timer(1.5, _open_browser).start()
    app.run(host="127.0.0.1", port=5000, debug=False)
