"""
Extract performance data from monthly Excel files (YYYY年M月一线考核.xlsx)
and write a consolidated JSON cache.
"""

import json
import os
import re
from datetime import datetime

from openpyxl import load_workbook

FILE_PATTERN = re.compile(r"^(\d{4})年(\d{1,2})月一线考核.*\.xlsx$")

# Normalise names that contain extra whitespace (e.g. "陈  杰" → "陈杰")
_WHITESPACE = re.compile(r"\s+")

# Map variant office names to a canonical form
_OFFICE_ALIASES = {
    "高新区分公司": "高新分公司",
}


def _clean_name(raw):
    if raw is None:
        return None
    return _WHITESPACE.sub("", str(raw).strip()) or None


def _normalise_office(raw):
    if raw is None:
        return None
    name = _WHITESPACE.sub("", str(raw).strip())
    return _OFFICE_ALIASES.get(name, name) or None


def _to_float(val):
    try:
        return round(float(val), 2)
    except (TypeError, ValueError):
        return None


def _clean_result(val):
    if val is None:
        return None
    s = str(val).strip().upper()
    return s if s in ("A", "B", "C") else None


# ---------------------------------------------------------------------------
# Per-sheet extraction
# ---------------------------------------------------------------------------

def _extract_dept_eval(wb):
    """部门绩效考核: office-level marks/results (rows 3/5/6, cols B+)."""
    ws = wb["部门绩效考核"]
    records = []
    for col in range(2, ws.max_column + 1):
        office = _normalise_office(ws.cell(row=3, column=col).value)
        if not office:
            continue
        mark = _to_float(ws.cell(row=5, column=col).value)
        result = _clean_result(ws.cell(row=6, column=col).value)
        if mark is not None or result is not None:
            records.append({"office": office, "mark": mark, "result": result})
    return records


def _extract_mgmt(wb):
    """部门副职考核: manager names/marks/results (rows 3/4/9/12, cols B+)."""
    ws = wb["部门副职考核"]
    records = []
    for col in range(2, ws.max_column + 1):
        office = _normalise_office(ws.cell(row=3, column=col).value)
        name = _clean_name(ws.cell(row=4, column=col).value)
        if not office or not name:
            continue
        mark = _to_float(ws.cell(row=9, column=col).value)
        result = _clean_result(ws.cell(row=12, column=col).value)
        records.append({
            "name": name,
            "office": office,
            "mark": mark,
            "result": result,
        })
    return records


def _extract_employees(wb):
    """绩效考评汇总表: employee names/marks/results with office sections."""
    ws = wb["绩效考评汇总表"]
    records = []
    current_office = None

    for row in range(4, ws.max_row + 1):
        val_a = ws.cell(row=row, column=1).value
        if val_a:
            text = str(val_a).strip()
            if text.startswith("备注"):
                break
            current_office = _normalise_office(text)
            if not current_office:
                continue

        name = _clean_name(ws.cell(row=row, column=7).value)
        if not name or not current_office:
            continue

        mark = _to_float(ws.cell(row=row, column=9).value)
        result = _clean_result(ws.cell(row=row, column=10).value)
        records.append({
            "name": name,
            "office": current_office,
            "mark": mark,
            "result": result,
        })
    return records


# ---------------------------------------------------------------------------
# Main extraction
# ---------------------------------------------------------------------------

def extract_all(data_dir):
    """Read every matching Excel file in *data_dir* and return a dict ready
    for JSON serialisation."""

    office_scores = {}   # office -> [{year, month, mark, result}]
    people = {}          # person_name -> [{year, month, office, type, mark, result}]
    files_processed = []

    for fname in sorted(os.listdir(data_dir)):
        if fname.startswith("~$"):
            continue
        m = FILE_PATTERN.match(fname)
        if not m:
            continue

        year, month = int(m.group(1)), int(m.group(2))
        fpath = os.path.join(data_dir, fname)

        try:
            wb = load_workbook(fpath, data_only=True)
        except Exception as exc:
            print(f"[WARN] Cannot open {fname}: {exc}")
            continue

        try:
            # Department-level scores
            for rec in _extract_dept_eval(wb):
                office_scores.setdefault(rec["office"], []).append({
                    "year": year, "month": month,
                    "mark": rec["mark"], "result": rec["result"],
                })

            # Managers
            for rec in _extract_mgmt(wb):
                people.setdefault(rec["name"], []).append({
                    "year": year, "month": month,
                    "office": rec["office"],
                    "type": "manager",
                    "mark": rec["mark"], "result": rec["result"],
                })

            # Employees
            for rec in _extract_employees(wb):
                people.setdefault(rec["name"], []).append({
                    "year": year, "month": month,
                    "office": rec["office"],
                    "type": "employee",
                    "mark": rec["mark"], "result": rec["result"],
                })

            files_processed.append(fname)
        finally:
            wb.close()

    return {
        "last_updated": datetime.now().isoformat(timespec="seconds"),
        "files_processed": files_processed,
        "office_scores": office_scores,
        "people": people,
    }


def run(data_dir, out_path):
    data = extract_all(data_dir)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"Extracted {len(data['files_processed'])} file(s), "
          f"{len(data['people'])} people, "
          f"{len(data['office_scores'])} offices → {out_path}")
    return data


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    data_dir = os.path.join(here, "..", "data")
    out_path = os.path.join(here, "performance_data.json")
    run(data_dir, out_path)
