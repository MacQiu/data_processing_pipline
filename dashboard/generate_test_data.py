"""
Generate test Excel files from the 2026年2月 template with randomized
marks, results, and occasional personnel transfers across offices.

Usage:  python generate_test_data.py
Output: data/YYYY年M月一线考核.xlsx  (for each month in the range)
"""

import os
import random
import re
import shutil

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(HERE, "..", "data")
TEMPLATE = os.path.join(DATA_DIR, "2026年2月一线考核.xlsx")

MONTHS = []
for y in range(2025, 2027):
    for m in range(1, 13):
        if y == 2025 and m < 1:
            continue
        if y == 2026 and m > 6:
            break
        MONTHS.append((y, m))

SKIP = {(2026, 2)}


def _jitter(base, spread=3.0):
    return round(base + random.uniform(-spread, spread), 2)


def _read_template_data():
    """Read original values from the template so we have base marks to
    randomize around."""
    wb = load_workbook(TEMPLATE, data_only=True)

    # Sheet 1: dept eval
    ws1 = wb["部门绩效考核"]
    dept_offices = []
    dept_marks = []
    for col in range(2, ws1.max_column + 1):
        name = ws1.cell(row=3, column=col).value
        mark = ws1.cell(row=5, column=col).value
        if name and mark is not None:
            dept_offices.append((col, str(name).strip()))
            dept_marks.append(float(mark))

    # Sheet 2: mgmt
    ws2 = wb["部门副职考核"]
    mgmt_cols = []
    mgmt_marks = []
    mgmt_names = []
    for col in range(2, ws2.max_column + 1):
        office = ws2.cell(row=3, column=col).value
        name = ws2.cell(row=4, column=col).value
        mark = ws2.cell(row=9, column=col).value
        if office and name and mark is not None:
            mgmt_cols.append(col)
            mgmt_marks.append(float(mark))
            mgmt_names.append(str(name).strip())

    # Sheet 3: employees -- collect (row, base_mark) per office section
    ws3 = wb["绩效考评汇总表"]
    emp_sections = []
    current_office = None
    current_rows = []
    for row in range(4, ws3.max_row + 1):
        val_a = ws3.cell(row=row, column=1).value
        if val_a:
            text = str(val_a).strip()
            if text.startswith("备注"):
                if current_office and current_rows:
                    emp_sections.append((current_office, list(current_rows)))
                break
            if current_office and current_rows:
                emp_sections.append((current_office, list(current_rows)))
            current_office = text
            current_rows = []
        name = ws3.cell(row=row, column=7).value
        mark = ws3.cell(row=row, column=9).value
        if name and mark is not None:
            current_rows.append((row, float(mark)))

    wb.close()
    return {
        "dept_offices": dept_offices,
        "dept_marks": dept_marks,
        "mgmt_cols": mgmt_cols,
        "mgmt_marks": mgmt_marks,
        "mgmt_names": list(mgmt_names),
        "emp_sections": emp_sections,
    }


def _assign_dept_results(marks_with_idx):
    """Office-level results: top mark = A, rest = B, below 75 = C."""
    if not marks_with_idx:
        return {}
    sorted_list = sorted(marks_with_idx, key=lambda x: x[1], reverse=True)
    results = {}
    for i, (idx, mark) in enumerate(sorted_list):
        if i == 0:
            results[idx] = "A"
        elif mark < 75:
            results[idx] = "C"
        else:
            results[idx] = "B"
    return results


def _assign_people_results(marks_with_idx):
    """People results: top mark = A, everyone else = B (no C)."""
    if not marks_with_idx:
        return {}
    sorted_list = sorted(marks_with_idx, key=lambda x: x[1], reverse=True)
    results = {}
    for i, (idx, _mark) in enumerate(sorted_list):
        results[idx] = "A" if i == 0 else "B"
    return results


def _generate_one(year, month, tpl_data, out_dir=None):
    fname = f"{year}年{month}月一线考核.xlsx"
    dest = os.path.join(out_dir or DATA_DIR, fname)

    shutil.copy2(TEMPLATE, dest)
    wb = load_workbook(dest)

    # --- Sheet 1: 部门绩效考核 ---
    ws1 = wb["部门绩效考核"]
    ws1.cell(row=1, column=1).value = f"{year}年{month}月部门考核评定汇总表"

    new_dept_marks = []
    for i, (col, _name) in enumerate(tpl_data["dept_offices"]):
        m = _jitter(tpl_data["dept_marks"][i], 3.0)
        m = max(60, min(100, m))
        new_dept_marks.append((col, m))
        ws1.cell(row=5, column=col).value = round(m, 2)

    sorted_depts = sorted(new_dept_marks, key=lambda x: x[1], reverse=True)
    dept_results = _assign_dept_results(new_dept_marks)
    for rank, (col, _mark) in enumerate(sorted_depts, 1):
        ws1.cell(row=4, column=col).value = rank
    for col, _mark in new_dept_marks:
        ws1.cell(row=6, column=col).value = dept_results[col]

    # --- Sheet 2: 部门副职考核 ---
    ws2 = wb["部门副职考核"]
    ws2.cell(row=1, column=1).value = f"{year}年{month}月中层副职管理人员考核评定汇总表"

    mgmt_names = list(tpl_data["mgmt_names"])
    if random.random() < 0.12 and len(mgmt_names) >= 2:
        i, j = random.sample(range(len(mgmt_names)), 2)
        mgmt_names[i], mgmt_names[j] = mgmt_names[j], mgmt_names[i]
        for idx, col in enumerate(tpl_data["mgmt_cols"]):
            ws2.cell(row=4, column=col).value = mgmt_names[idx]

    new_mgmt_marks = []
    for i, col in enumerate(tpl_data["mgmt_cols"]):
        m = _jitter(tpl_data["mgmt_marks"][i], 4.0)
        m = max(60, min(100, m))
        new_mgmt_marks.append((col, m))
        ws2.cell(row=9, column=col).value = round(m, 2)

    mgmt_results = _assign_people_results(new_mgmt_marks)
    for col, _mark in new_mgmt_marks:
        ws2.cell(row=12, column=col).value = mgmt_results[col]

    # --- Sheet 3: 绩效考评汇总表 ---
    ws3 = wb["绩效考评汇总表"]
    ws3.cell(row=1, column=1).value = f"{year}年{month}月绩效考核评定汇总表"
    ws3.cell(row=3, column=10).value = f"{month}月员工\n综合评定结果"

    all_emp_rows = []
    for _office, rows in tpl_data["emp_sections"]:
        section_marks = []
        for row_num, base_mark in rows:
            m = _jitter(base_mark, 3.0)
            m = max(60, min(100, m))
            section_marks.append((row_num, m))
            ws3.cell(row=row_num, column=9).value = round(m, 2)
        section_results = _assign_people_results(section_marks)
        for row_num, _mark in section_marks:
            ws3.cell(row=row_num, column=10).value = section_results[row_num]

    # Occasional employee transfer: swap two employees between different offices
    if random.random() < 0.10 and len(tpl_data["emp_sections"]) >= 2:
        sec_indices = random.sample(range(len(tpl_data["emp_sections"])), 2)
        sec_a = tpl_data["emp_sections"][sec_indices[0]]
        sec_b = tpl_data["emp_sections"][sec_indices[1]]
        if sec_a[1] and sec_b[1]:
            row_a = random.choice(sec_a[1])[0]
            row_b = random.choice(sec_b[1])[0]
            name_a = ws3.cell(row=row_a, column=7).value
            name_b = ws3.cell(row=row_b, column=7).value
            ws3.cell(row=row_a, column=7).value = name_b
            ws3.cell(row=row_b, column=7).value = name_a

    wb.save(dest)
    wb.close()
    return fname


def main():
    if not os.path.exists(TEMPLATE):
        print(f"Template not found: {TEMPLATE}")
        return

    tpl_data = _read_template_data()
    generated = []

    for year, month in MONTHS:
        if (year, month) in SKIP:
            print(f"  skip {year}-{month:02d} (real data)")
            continue
        fname = _generate_one(year, month, tpl_data)
        generated.append(fname)
        print(f"  {fname}")

    print(f"\nGenerated {len(generated)} test files in {DATA_DIR}")


def generate_upload_test():
    """Generate a single test file outside data/ for upload testing."""
    if not os.path.exists(TEMPLATE):
        print(f"Template not found: {TEMPLATE}")
        return
    out_dir = os.path.join(HERE, "test_upload")
    os.makedirs(out_dir, exist_ok=True)
    tpl_data = _read_template_data()
    fname = _generate_one(2026, 7, tpl_data, out_dir=out_dir)
    print(f"Test upload file: {os.path.join(out_dir, fname)}")


if __name__ == "__main__":
    random.seed(42)
    main()
    generate_upload_test()
